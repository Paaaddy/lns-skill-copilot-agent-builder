#!/usr/bin/env python3
"""
generate_docx.py — Word document generator for Microsoft 365 Copilot agents.
Trilingual labels: English ("en"), French ("fr") and German ("de"), chosen via
the top-level `lang` field in the JSON config. Defaults to "en" when unset.

Usage: python generate_docx.py <config.json> <output.docx>
"""

import json
import sys
from pathlib import Path
from datetime import datetime

try:
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
except ImportError:
    print("Installing python-docx...")
    import subprocess
    subprocess.run([sys.executable, "-m", "pip", "install", "python-docx", "--user", "-q"])
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement


# ── Palette ────────────────────────────────────────────────────────────────────
ORANGE      = RGBColor(0xFF, 0x51, 0x19)
DARK        = RGBColor(0x1A, 0x1A, 0x1A)
GRAY        = RGBColor(0x88, 0x88, 0x88)
LIGHT_GRAY  = RGBColor(0xF5, 0xF5, 0xF5)
RULE_COLOR  = "DDDDDD"

# ── Typography ─────────────────────────────────────────────────────────────────
FONT_TITLE  = "Calibri"
FONT_BODY   = "Georgia"
SIZE_HERO   = 28
SIZE_H1     = 16
SIZE_H2     = 12
SIZE_BODY   = 11
SIZE_META   = 9
SIZE_SMALL  = 8

VERSION     = "2.0.0"

# ── Capability labels per language ──────────────────────────────────────────────
CAPABILITIES_LABELS = {
    "en": {
        "WebSearch":             "🔍 Web search",
        "OneDriveAndSharePoint": "📁 OneDrive & SharePoint",
        "Email":                 "📧 Email",
        "TeamsMessages":         "💬 Teams messages",
        "People":                "👥 People",
        "Meetings":              "📅 Meetings",
        "GraphicArt":            "🎨 Image creation",
        "CodeInterpreter":       "💻 Code interpreter",
        "Dataverse":             "🗄️ Dataverse",
    },
    "fr": {
        "WebSearch":             "🔍 Recherche web",
        "OneDriveAndSharePoint": "📁 OneDrive & SharePoint",
        "Email":                 "📧 E-mail",
        "TeamsMessages":         "💬 Messages Teams",
        "People":                "👥 Personnes",
        "Meetings":              "📅 Réunions",
        "GraphicArt":            "🎨 Création d'images",
        "CodeInterpreter":       "💻 Interpréteur de code",
        "Dataverse":             "🗄️ Dataverse",
    },
    "de": {
        "WebSearch":             "🔍 Websuche",
        "OneDriveAndSharePoint": "📁 OneDrive & SharePoint",
        "Email":                 "📧 E-Mail",
        "TeamsMessages":         "💬 Teams-Nachrichten",
        "People":                "👥 Personen",
        "Meetings":              "📅 Besprechungen",
        "GraphicArt":            "🎨 Bilderstellung",
        "CodeInterpreter":       "💻 Code-Interpreter",
        "Dataverse":             "🗄️ Dataverse",
    },
}

# ── UI text per language ────────────────────────────────────────────────────────
LABELS = {
    "fr": {
        "default_agent_name":   "Mon agent Copilot",
        "date_format":          "%d/%m/%Y",
        "supertitle":           "MICROSOFT 365 COPILOT — AGENT BUILDER",
        "subtitle":             "Configuration complète — prête à coller dans Agent Builder",
        "generated_on":         "Généré le {date}   ·   https://m365.cloud.microsoft/chat/agent/new",
        "not_set":              "(non renseigné)",
        "char_unit":            "caractères",
        "section_identity":     "Identité de l'agent",
        "section_instructions": "Instructions",
        "section_capabilities": "Fonctionnalités à activer",
        "section_sources":      "Sources de connaissances",
        "section_starters":     "Suggestions de démarrage",
        "section_disclaimer":   "Disclaimer",
        "section_next_steps":   "Prochaines étapes recommandées",
        "field_name":           "Nom",
        "field_description":    "Description",
        "field_instructions":   "Instructions",
        "field_capabilities":   "Fonctionnalités",
        "field_sources":        "Sources (URLs SharePoint / sites web)",
        "field_starters":       "Suggestions de démarrage",
        "field_disclaimer":     "Message affiché au démarrage",
        "hint_name":            "Copiez ce nom dans le champ « Nom » de l'Agent Builder.",
        "hint_description":     "Cette description est visible par les utilisateurs qui parcourent la liste des agents.",
        "hint_instructions":    "Le champ le plus important. Collez l'intégralité du contenu dans le champ « Instructions ». Affinez après vos premiers tests.",
        "hint_capabilities":    "Activez ces fonctionnalités dans l'onglet « Fonctionnalités » de l'Agent Builder.",
        "hint_sources":         "Ajoutez ces sources dans l'onglet « Connaissances ». Vérifiez que vos utilisateurs ont accès aux documents sous-jacents.",
        "hint_disclaimer":      "Ce texte est affiché à l'utilisateur au début de chaque conversation.",
        "no_capabilities":      "(aucune fonctionnalité spécifiée)",
        "starter_count":        "{n} suggestions de démarrage configurées   ·   maximum 12",
        "checklist_title":      "CHECKLIST DE MISE EN LIGNE",
        "checklist_items": [
            ("Ouvrir l'Agent Builder",              "https://m365.cloud.microsoft/chat/agent/new"),
            ("Renseigner le Nom et la Description", "Champs 1 et 2"),
            ("Coller les Instructions",             "Le champ le plus important — coller en intégralité"),
            ("Activer les Fonctionnalités",         "Onglet « Fonctionnalités »"),
            ("Ajouter les Sources de connaissances","SharePoint / OneDrive — onglet « Connaissances »"),
            ("Saisir les Suggestions de démarrage", "Jusqu'à 12 suggestions"),
            ("Tester en mode Aperçu",               "Lancer quelques suggestions, affiner si nécessaire"),
            ("Publier et partager",                 "Diffuser auprès des utilisateurs cibles"),
        ],
        "footer":               "Document généré le {date}   ·   Microsoft 365 Copilot Agent Builder   ·   v{version}",
        "success":              "✅ Document généré : {path}",
    },
    "en": {
        "default_agent_name":   "My Copilot Agent",
        "date_format":          "%Y-%m-%d",
        "supertitle":           "MICROSOFT 365 COPILOT — AGENT BUILDER",
        "subtitle":             "Complete configuration — ready to paste into Agent Builder",
        "generated_on":         "Generated on {date}   ·   https://m365.cloud.microsoft/chat/agent/new",
        "not_set":              "(not set)",
        "char_unit":            "chars",
        "section_identity":     "Agent identity",
        "section_instructions": "Instructions",
        "section_capabilities": "Capabilities to enable",
        "section_sources":      "Knowledge sources",
        "section_starters":     "Starter prompts",
        "section_disclaimer":   "Disclaimer",
        "section_next_steps":   "Recommended next steps",
        "field_name":           "Name",
        "field_description":    "Description",
        "field_instructions":   "Instructions",
        "field_capabilities":   "Capabilities",
        "field_sources":        "Sources (SharePoint URLs / websites)",
        "field_starters":       "Starter prompts",
        "field_disclaimer":     "Message shown at start",
        "hint_name":            "Copy this name into the \"Name\" field of the Agent Builder.",
        "hint_description":     "This description is shown to users browsing the agent list.",
        "hint_instructions":    "The most important field. Paste the whole content into the \"Instructions\" field. Refine after your first tests.",
        "hint_capabilities":    "Enable these capabilities in the \"Capabilities\" tab of the Agent Builder.",
        "hint_sources":         "Add these sources in the \"Knowledge\" tab. Make sure your users have access to the underlying documents.",
        "hint_disclaimer":      "This text is shown to the user at the start of every conversation.",
        "no_capabilities":      "(no capabilities specified)",
        "starter_count":        "{n} starter prompts configured   ·   max 12",
        "checklist_title":      "GO-LIVE CHECKLIST",
        "checklist_items": [
            ("Open the Agent Builder",         "https://m365.cloud.microsoft/chat/agent/new"),
            ("Fill in Name and Description",   "Fields 1 and 2"),
            ("Paste the Instructions",         "The most important field — paste in full"),
            ("Enable the Capabilities",        "\"Capabilities\" tab"),
            ("Add the Knowledge sources",      "SharePoint / OneDrive — \"Knowledge\" tab"),
            ("Enter the Starter prompts",      "Up to 12 starter prompts"),
            ("Test in Preview mode",           "Run a few starters, refine as needed"),
            ("Publish and share",              "Share with your target users"),
        ],
        "footer":               "Document generated on {date}   ·   Microsoft 365 Copilot Agent Builder   ·   v{version}",
        "success":              "✅ Document generated: {path}",
    },
    "de": {
        "default_agent_name":   "Mein Copilot-Agent",
        "date_format":          "%d.%m.%Y",
        "supertitle":           "MICROSOFT 365 COPILOT — AGENT BUILDER",
        "subtitle":             "Vollständige Konfiguration — bereit zum Einfügen in den Agent Builder",
        "generated_on":         "Erstellt am {date}   ·   https://m365.cloud.microsoft/chat/agent/new",
        "not_set":              "(nicht angegeben)",
        "char_unit":            "Zeichen",
        "section_identity":     "Identität des Agenten",
        "section_instructions": "Anweisungen",
        "section_capabilities": "Zu aktivierende Funktionen",
        "section_sources":      "Wissensquellen",
        "section_starters":     "Startvorschläge",
        "section_disclaimer":   "Disclaimer",
        "section_next_steps":   "Empfohlene nächste Schritte",
        "field_name":           "Name",
        "field_description":    "Beschreibung",
        "field_instructions":   "Anweisungen",
        "field_capabilities":   "Funktionen",
        "field_sources":        "Quellen (SharePoint-URLs / Websites)",
        "field_starters":       "Startvorschläge",
        "field_disclaimer":     "Nachricht zu Beginn",
        "hint_name":            "Kopiere diesen Namen in das Feld „Name“ des Agent Builders.",
        "hint_description":     "Diese Beschreibung sehen die Nutzer:innen in der Agentenliste.",
        "hint_instructions":    "Das wichtigste Feld. Füge den gesamten Inhalt in das Feld „Anweisungen“ ein. Nach den ersten Tests feinjustieren.",
        "hint_capabilities":    "Aktiviere diese Funktionen im Tab „Funktionen“ des Agent Builders.",
        "hint_sources":         "Füge diese Quellen im Tab „Wissen“ hinzu. Stelle sicher, dass die Nutzer:innen Zugriff auf die Dokumente haben.",
        "hint_disclaimer":      "Dieser Text wird den Nutzer:innen zu Beginn jeder Unterhaltung angezeigt.",
        "no_capabilities":      "(keine Funktionen angegeben)",
        "starter_count":        "{n} Startvorschläge konfiguriert   ·   Maximum 12",
        "checklist_title":      "GO-LIVE-CHECKLISTE",
        "checklist_items": [
            ("Agent Builder öffnen",           "https://m365.cloud.microsoft/chat/agent/new"),
            ("Name und Beschreibung eintragen","Felder 1 und 2"),
            ("Anweisungen einfügen",           "Wichtigstes Feld — vollständig einfügen"),
            ("Funktionen aktivieren",          "Tab „Funktionen“"),
            ("Quellen hinzufügen",             "SharePoint / OneDrive — Tab „Wissen“"),
            ("Startvorschläge eintragen",      "Bis zu 12 Startvorschläge"),
            ("Im Preview-Modus testen",        "Startvorschläge ausprobieren, bei Bedarf anpassen"),
            ("Veröffentlichen und teilen",     "An die Zielnutzer:innen verteilen"),
        ],
        "footer":               "Dokument erstellt am {date}   ·   Microsoft 365 Copilot Agent Builder   ·   v{version}",
        "success":              "✅ Dokument erstellt: {path}",
    },
}


# ── Next-step recommendations per agent type ────────────────────────────────────
NEXT_STEPS = {
    "en": {
        "hr": [
            "Upload your HR policies, employee handbook, and org chart to SharePoint.",
            "Grant the agent access to the relevant SharePoint folder via the OneDrive & SharePoint capability.",
            "Run a pilot with 5–10 HR team members before rolling out organisation-wide.",
            "Schedule a monthly review of the instructions as policies evolve.",
        ],
        "it": [
            "Centralise system documentation and technical FAQs in a shared SharePoint library.",
            "Define escalation paths explicitly in the instructions (Tier 1 → Tier 2 → human).",
            "Test edge cases: unknown errors, out-of-scope requests, sensitive data handling.",
            "Monitor usage monthly to identify gaps in the knowledge base.",
        ],
        "sales": [
            "Upload your pitch deck, price list, and objection FAQ to SharePoint.",
            "Update knowledge sources quarterly as product catalogue and pricing change.",
            "Include 2–3 real deal examples in the knowledge base to anchor the tone.",
            "Brief the sales team with a 15-minute onboarding session before launch.",
        ],
        "legal": [
            "Ensure all uploaded templates are the current approved versions.",
            "Add an explicit out-of-scope rule: the agent provides information, not legal advice.",
            "Have the legal team validate the instructions before going live.",
            "Plan a quarterly review cycle aligned with regulatory update schedules.",
        ],
        "pm": [
            "Upload your methodology templates and project glossary to SharePoint.",
            "Align the starter prompts with your team's most common recurring tasks.",
            "Test the agent on a live project before rolling out across all teams.",
            "Encode escalation paths for blocked decisions or resource conflicts.",
        ],
        "writing": [
            "Upload top-performing past articles and your editorial style guide.",
            "Include 3–5 annotated examples showing the tone and structure you want.",
            "Run the anti-pattern pass on the agent's first 10 outputs and refine the instructions.",
            "Create a shared SharePoint folder for approved content samples to grow over time.",
        ],
        "data": [
            "Enable the Code Interpreter capability for data analysis tasks.",
            "Define the expected output format (tables, charts, summaries) in the instructions.",
            "Upload data dictionaries and schema documentation to the knowledge base.",
            "Add explicit instructions for handling missing data, outliers, and confidentiality.",
        ],
        "custom": [
            "Review the instructions with a sample of target users before going live.",
            "Identify the 3 most common user requests and optimise the starter prompts for them.",
            "Schedule a 30-day checkpoint to review usage and refine the instructions.",
            "Document the agent's scope and limitations in a shared team page.",
        ],
    },
    "fr": {
        "hr": [
            "Chargez vos politiques RH, votre règlement intérieur et votre organigramme sur SharePoint.",
            "Accordez à l'agent l'accès au dossier SharePoint concerné via la fonctionnalité OneDrive & SharePoint.",
            "Effectuez un pilote avec 5 à 10 membres de l'équipe RH avant le déploiement à l'échelle.",
            "Planifiez une révision mensuelle des instructions à mesure que les politiques évoluent.",
        ],
        "it": [
            "Centralisez la documentation système et les FAQ techniques dans une bibliothèque SharePoint partagée.",
            "Définissez les chemins d'escalade explicitement dans les instructions (Niveau 1 → Niveau 2 → humain).",
            "Testez les cas limites : erreurs inconnues, demandes hors périmètre, gestion des données sensibles.",
            "Consultez les journaux d'utilisation mensuellement pour identifier les lacunes dans la base de connaissances.",
        ],
        "sales": [
            "Chargez votre pitch deck, liste de prix et FAQ objections sur SharePoint.",
            "Mettez à jour les sources trimestriellement au rythme des évolutions produit et tarifaires.",
            "Incluez 2 à 3 exemples réels de deals dans la base de connaissances pour ancrer le ton.",
            "Briefez l'équipe commerciale avec une session d'onboarding de 15 minutes avant le lancement.",
        ],
        "legal": [
            "Vérifiez que tous les modèles chargés sont bien les versions actuellement approuvées.",
            "Ajoutez une règle hors périmètre explicite : l'agent donne des informations, pas des conseils juridiques.",
            "Faites valider les instructions par l'équipe juridique avant la mise en production.",
            "Planifiez un cycle de révision trimestriel aligné sur les calendriers de mise à jour réglementaire.",
        ],
        "pm": [
            "Chargez vos modèles de méthodologie et votre glossaire projet sur SharePoint.",
            "Alignez les suggestions de démarrage sur les tâches récurrentes les plus fréquentes de votre équipe.",
            "Testez l'agent sur un projet en cours avant de le déployer à l'ensemble des équipes.",
            "Encodez les chemins d'escalade pour les décisions bloquées ou les conflits de ressources.",
        ],
        "writing": [
            "Chargez vos articles les plus performants et votre guide de style éditorial.",
            "Incluez 3 à 5 exemples annotés montrant le ton et la structure souhaités.",
            "Effectuez le passage anti-patterns sur les 10 premières productions de l'agent et affinez les instructions.",
            "Créez un dossier SharePoint partagé pour les exemples de contenus validés, à enrichir au fil du temps.",
        ],
        "data": [
            "Activez la fonctionnalité Code Interpreter pour les tâches d'analyse de données.",
            "Définissez le format de sortie attendu (tableaux, graphiques, synthèses) dans les instructions.",
            "Chargez les dictionnaires de données et la documentation des schémas dans la base de connaissances.",
            "Ajoutez des instructions explicites pour gérer les données manquantes, les valeurs aberrantes et la confidentialité.",
        ],
        "custom": [
            "Révisez les instructions avec un échantillon d'utilisateurs cibles avant le lancement.",
            "Identifiez les 3 demandes les plus fréquentes et optimisez les suggestions de démarrage en conséquence.",
            "Planifiez un point de contrôle à 30 jours pour analyser les usages et affiner les instructions.",
            "Documentez le périmètre et les limites de l'agent dans une page d'équipe partagée.",
        ],
    },
    "de": {
        "hr": [
            "Lade deine HR-Richtlinien, die Arbeitsordnung und das Organigramm auf SharePoint hoch.",
            "Erteile dem Agenten Zugriff auf den relevanten SharePoint-Ordner via OneDrive & SharePoint.",
            "Führe einen Pilottest mit 5–10 HR-Teammitgliedern durch, bevor du organisationsweit ausrollst.",
            "Plane eine monatliche Überprüfung der Anweisungen, während sich Richtlinien weiterentwickeln.",
        ],
        "it": [
            "Zentralisiere die Systemdokumentation und technische FAQs in einer gemeinsamen SharePoint-Bibliothek.",
            "Definiere Eskalationspfade explizit in den Anweisungen (Stufe 1 → Stufe 2 → Mensch).",
            "Teste Grenzfälle: unbekannte Fehler, Anfragen außerhalb des Bereichs, Umgang mit sensiblen Daten.",
            "Überprüfe monatlich die Nutzungsprotokolle, um Lücken in der Wissensbasis zu identifizieren.",
        ],
        "sales": [
            "Lade dein Pitch Deck, die Preisliste und die Einwände-FAQ auf SharePoint hoch.",
            "Aktualisiere die Wissensquellen vierteljährlich entsprechend der Produkt- und Preisänderungen.",
            "Füge 2–3 echte Deal-Beispiele in die Wissensbasis ein, um den Ton zu verankern.",
            "Briefiere das Vertriebsteam mit einer 15-minütigen Onboarding-Session vor dem Start.",
        ],
        "legal": [
            "Stelle sicher, dass alle hochgeladenen Vorlagen die aktuell genehmigten Versionen sind.",
            "Füge eine explizite Außer-Bereich-Regel ein: Der Agent gibt Informationen, keine Rechtsberatung.",
            "Koordiniere mit dem Rechtsteam, um die Anweisungen vor der Inbetriebnahme zu validieren.",
            "Plane einen vierteljährlichen Überprüfungszyklus abgestimmt auf regulatorische Update-Zeitpläne.",
        ],
        "pm": [
            "Lade deine Methodikvorlagen und das Projektglossar auf SharePoint hoch.",
            "Stimme die Startvorschläge auf die häufigsten wiederkehrenden Aufgaben deines Teams ab.",
            "Teste den Agenten mit einem laufenden Projekt, bevor du ihn in allen Teams ausrollst.",
            "Kodiere Eskalationspfade für blockierte Entscheidungen oder Ressourcenkonflikte.",
        ],
        "writing": [
            "Lade deine erfolgreichsten Artikel und deinen redaktionellen Style Guide hoch.",
            "Füge 3–5 kommentierte Beispiele ein, die Ton und Struktur zeigen.",
            "Führe den Anti-Muster-Durchgang bei den ersten 10 Ausgaben des Agenten durch und verfeinere die Anweisungen.",
            "Richte einen gemeinsamen SharePoint-Ordner für freigegebene Inhaltsbeispiele ein.",
        ],
        "data": [
            "Aktiviere die Funktion Code Interpreter für Datenanalyseaufgaben.",
            "Definiere das erwartete Ausgabeformat (Tabellen, Diagramme, Zusammenfassungen) in den Anweisungen.",
            "Lade Datenwörterbücher und Schema-Dokumentation in die Wissensbasis hoch.",
            "Füge explizite Anweisungen für den Umgang mit fehlenden Daten, Ausreißern und Vertraulichkeit hinzu.",
        ],
        "custom": [
            "Überprüfe die Anweisungen mit einer Auswahl von Zielnutzern vor dem Launch.",
            "Identifiziere die 3 häufigsten Nutzeranfragen und optimiere die Startvorschläge dafür.",
            "Plane einen 30-Tage-Kontrollpunkt, um Nutzung zu überprüfen und Anweisungen zu verfeinern.",
            "Dokumentiere den Bereich und die Grenzen des Agenten auf einer gemeinsamen Team-Seite.",
        ],
    },
}


# ── XML helpers ─────────────────────────────────────────────────────────────────

# Pre-compute qualified names to avoid repeated qn() calls (micro-optimization)
_QN_RFONTS   = qn('w:rFonts')
_QN_ASCII    = qn('w:ascii')
_QN_HANSI    = qn('w:hAnsi')
_QN_EASTASIA = qn('w:eastAsia')
_QN_CS       = qn('w:cs')
_QN_VAL      = qn('w:val')
_QN_SZ       = qn('w:sz')
_QN_SPACE    = qn('w:space')
_QN_COLOR    = qn('w:color')
_QN_PBDR     = qn('w:pBdr')
_QN_BOTTOM   = qn('w:bottom')
_QN_LEFT     = qn('w:left')
_QN_SHD      = qn('w:shd')
_QN_FILL     = qn('w:fill')

def set_run_font(run, font_name, size_pt, bold=False, color=None, italic=False):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()

    existing = rPr.find(_QN_RFONTS)
    if existing is not None:
        rPr.remove(existing)

    rFonts = OxmlElement('w:rFonts')
    rFonts.set(_QN_ASCII,    font_name)
    rFonts.set(_QN_HANSI,    font_name)
    rFonts.set(_QN_EASTASIA, font_name)
    rFonts.set(_QN_CS,       font_name)
    rPr.insert(0, rFonts)


def add_rule(doc, before_pt=6, after_pt=12, color=RULE_COLOR):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before_pt)
    p.paragraph_format.space_after  = Pt(after_pt)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(_QN_VAL,   'single')
    bottom.set(_QN_SZ,    '4')
    bottom.set(_QN_SPACE, '1')
    bottom.set(_QN_COLOR, color)
    pBdr.append(bottom)
    pPr.append(pBdr)
    return p


def set_para_shading(para, fill_hex):
    pPr = para._p.get_or_add_pPr()
    existing = pPr.find(_QN_SHD)
    if existing is not None:
        pPr.remove(existing)
    shd = OxmlElement('w:shd')
    shd.set(_QN_VAL,   'clear')
    shd.set(_QN_COLOR, 'auto')
    shd.set(_QN_FILL,  fill_hex)
    pPr.append(shd)


def add_left_border(para, color_hex="FF5119", size="12"):
    pPr = para._p.get_or_add_pPr()
    pBdr = pPr.find(_QN_PBDR)
    if pBdr is None:
        pBdr = OxmlElement('w:pBdr')
        pPr.append(pBdr)
    existing_left = pBdr.find(_QN_LEFT)
    if existing_left is not None:
        pBdr.remove(existing_left)
    left = OxmlElement('w:left')
    left.set(_QN_VAL,   'single')
    left.set(_QN_SZ,    size)
    left.set(_QN_SPACE, '12')
    left.set(_QN_COLOR, color_hex)
    pBdr.append(left)


# ── Content components ──────────────────────────────────────────────────────────

def add_section_title(doc, text, emoji=""):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(20)
    p.paragraph_format.space_after  = Pt(4)
    label = f"{emoji}  {text.upper()}" if emoji else text.upper()
    run = p.add_run(label)
    set_run_font(run, FONT_TITLE, SIZE_H1, bold=True, color=ORANGE)

    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(_QN_VAL,   'single')
    bottom.set(_QN_SZ,    '6')
    bottom.set(_QN_SPACE, '4')
    bottom.set(_QN_COLOR, 'FF5119')
    pBdr.append(bottom)
    pPr.append(pBdr)
    return p


def add_field_block(doc, field_name, content, char_limit=None, hint=None,
                    not_set_text="(not set)", char_unit="chars"):
    label_p = doc.add_paragraph()
    label_p.paragraph_format.space_before = Pt(12)
    label_p.paragraph_format.space_after  = Pt(3)

    label_run = label_p.add_run(field_name.upper())
    set_run_font(label_run, FONT_TITLE, SIZE_H2, bold=True, color=GRAY)

    if char_limit and content:
        count = len(content)
        sep = label_p.add_run(f"   {count} / {char_limit} {char_unit}")
        set_run_font(sep, FONT_TITLE, SIZE_META, color=GRAY)

    is_empty   = not content
    display_text = content or not_set_text
    lines      = display_text.split('\n')

    font_color        = GRAY if is_empty else DARK
    left_indent       = Cm(0.4)
    right_indent      = Cm(0.2)
    space_before_next = Pt(2)
    space_after       = Pt(2)

    for i, line in enumerate(lines):
        cp = doc.add_paragraph()
        cp.paragraph_format.space_before = space_before_next if i > 0 else Pt(0)
        cp.paragraph_format.space_after  = space_after
        cp.paragraph_format.left_indent  = left_indent
        cp.paragraph_format.right_indent = right_indent

        run = cp.add_run(line if line.strip() else " ")
        set_run_font(run, FONT_BODY, SIZE_BODY, color=font_color, italic=is_empty)
        set_para_shading(cp, "F5F5F5")
        add_left_border(cp)

    if hint:
        hp = doc.add_paragraph()
        hp.paragraph_format.space_before = Pt(3)
        hp.paragraph_format.space_after  = Pt(2)
        hp.paragraph_format.left_indent  = Cm(0.4)
        hr = hp.add_run(f"↗  {hint}")
        set_run_font(hr, FONT_TITLE, SIZE_META, color=GRAY, italic=True)


def add_suggestion_block(doc, index, title, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after  = Pt(2)

    num = p.add_run(f"{index}. ")
    set_run_font(num, FONT_TITLE, SIZE_BODY, bold=True, color=ORANGE)

    title_run = p.add_run(title)
    set_run_font(title_run, FONT_TITLE, SIZE_BODY, bold=True, color=DARK)

    tp = doc.add_paragraph()
    tp.paragraph_format.space_before = Pt(0)
    tp.paragraph_format.space_after  = Pt(4)
    tp.paragraph_format.left_indent  = Cm(0.5)
    tr = tp.add_run(f'"{text}"')
    set_run_font(tr, FONT_BODY, SIZE_BODY, color=GRAY, italic=True)


# ── Main generator ──────────────────────────────────────────────────────────────

def _get(config, *keys, default=""):
    """Return the first present non-empty value among keys in config."""
    return next((config[k] for k in keys if k in config and config[k] not in (None, "")), default)


def _validate_config(config: dict) -> None:
    """Raise ValueError for missing required fields."""
    name         = _get(config, "name", "nom")
    instructions = _get(config, "instructions")
    missing = []
    if not name:
        missing.append("name")
    if not instructions:
        missing.append("instructions")
    if missing:
        raise ValueError(f"Config missing required field(s): {', '.join(missing)}")
    lang = str(config.get("lang", "en")).lower()
    if lang not in LABELS:
        print(f"Warning: unknown lang '{lang}', defaulting to 'en'", file=sys.stderr)


def generate_document(config: dict, output_path: str) -> str:
    _validate_config(config)

    lang = str(config.get("lang", "en")).lower()
    if lang not in LABELS:
        lang = "en"
    L           = LABELS[lang]
    caps_labels = CAPABILITIES_LABELS[lang]

    # Resolve output path; avoid collisions with a microsecond timestamp suffix
    out = Path(output_path).resolve()
    if out.exists():
        ts  = datetime.now().strftime("%H%M%S%f")
        out = out.with_name(f"{out.stem}_{ts}{out.suffix}")

    doc = Document()

    for section in doc.sections:
        section.top_margin    = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin   = Cm(2.5)
        section.right_margin  = Cm(2.5)

    agent_name = _get(config, "name", "nom", default=L["default_agent_name"])
    date_str   = datetime.now().strftime(L["date_format"])

    # ══════════════════════════════════════════════════
    # COVER
    # ══════════════════════════════════════════════════

    sup = doc.add_paragraph()
    sup.paragraph_format.space_before = Pt(0)
    sup.paragraph_format.space_after  = Pt(6)
    sup.alignment = WD_ALIGN_PARAGRAPH.LEFT
    sup_run = sup.add_run(L["supertitle"])
    set_run_font(sup_run, FONT_TITLE, SIZE_META, bold=True, color=GRAY)

    hero = doc.add_paragraph()
    hero.paragraph_format.space_before = Pt(4)
    hero.paragraph_format.space_after  = Pt(8)
    hero.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hero_run = hero.add_run(agent_name)
    set_run_font(hero_run, FONT_TITLE, SIZE_HERO, bold=True, color=ORANGE)

    sub = doc.add_paragraph()
    sub.paragraph_format.space_before = Pt(0)
    sub.paragraph_format.space_after  = Pt(6)
    sub_run = sub.add_run(L["subtitle"])
    set_run_font(sub_run, FONT_BODY, SIZE_BODY, color=GRAY)

    meta = doc.add_paragraph()
    meta.paragraph_format.space_before = Pt(0)
    meta.paragraph_format.space_after  = Pt(16)
    meta_run = meta.add_run(L["generated_on"].format(date=date_str))
    set_run_font(meta_run, FONT_TITLE, SIZE_META, color=GRAY)

    add_rule(doc, before_pt=4, after_pt=20, color="FF5119")

    # ══════════════════════════════════════════════════
    # 1. IDENTITY
    # ══════════════════════════════════════════════════
    add_section_title(doc, L["section_identity"], "🤖")

    add_field_block(doc, L["field_name"], _get(config, "name", "nom"),
                    char_limit=100, hint=L["hint_name"],
                    not_set_text=L["not_set"], char_unit=L["char_unit"])

    add_field_block(doc, L["field_description"], _get(config, "description"),
                    char_limit=1000, hint=L["hint_description"],
                    not_set_text=L["not_set"], char_unit=L["char_unit"])

    # ══════════════════════════════════════════════════
    # 2. INSTRUCTIONS
    # ══════════════════════════════════════════════════
    add_section_title(doc, L["section_instructions"], "📋")

    add_field_block(doc, L["field_instructions"], _get(config, "instructions"),
                    char_limit=8000, hint=L["hint_instructions"],
                    not_set_text=L["not_set"], char_unit=L["char_unit"])

    # ══════════════════════════════════════════════════
    # 3. CAPABILITIES
    # ══════════════════════════════════════════════════
    add_section_title(doc, L["section_capabilities"], "⚙️")

    caps = _get(config, "capabilities", "fonctionnalites", default=[]) or []
    if caps:
        caps_text = "\n".join(f"✅  {caps_labels.get(c, c)}" for c in caps)
    else:
        caps_text = L["no_capabilities"]

    add_field_block(doc, L["field_capabilities"], caps_text,
                    hint=L["hint_capabilities"],
                    not_set_text=L["not_set"], char_unit=L["char_unit"])

    # ══════════════════════════════════════════════════
    # 4. KNOWLEDGE SOURCES
    # ══════════════════════════════════════════════════
    sources = _get(config, "knowledge_sources", "sources_connaissances", default=[]) or []
    if sources:
        add_section_title(doc, L["section_sources"], "📚")
        sources_text = "\n".join([f"•  {s}" for s in sources])
        add_field_block(doc, L["field_sources"], sources_text,
                        hint=L["hint_sources"],
                        not_set_text=L["not_set"], char_unit=L["char_unit"])

    # ══════════════════════════════════════════════════
    # 5. STARTER PROMPTS
    # ══════════════════════════════════════════════════
    starters = _get(config, "starters", "suggestions_demarrage", default=[]) or []
    if starters:
        add_section_title(doc, L["section_starters"], "💬")

        count_p = doc.add_paragraph()
        count_p.paragraph_format.space_before = Pt(2)
        count_p.paragraph_format.space_after  = Pt(8)
        count_run = count_p.add_run(L["starter_count"].format(n=len(starters)))
        set_run_font(count_run, FONT_TITLE, SIZE_META, color=GRAY)

        for i, s in enumerate(starters, 1):
            title = s.get("title") or s.get("titre") or f"{L['field_starters']} {i}"
            text  = s.get("text")  or s.get("texte")  or ""
            add_suggestion_block(doc, i, title, text)

    # ══════════════════════════════════════════════════
    # 6. DISCLAIMER
    # ══════════════════════════════════════════════════
    disclaimer = _get(config, "disclaimer")
    if disclaimer:
        add_section_title(doc, L["section_disclaimer"], "⚠️")
        add_field_block(doc, L["field_disclaimer"], disclaimer,
                        char_limit=500, hint=L["hint_disclaimer"],
                        not_set_text=L["not_set"], char_unit=L["char_unit"])

    # ══════════════════════════════════════════════════
    # NEXT STEPS
    # ══════════════════════════════════════════════════
    agent_type = str(config.get("agent_type", "custom")).lower()
    lang_steps = NEXT_STEPS.get(lang, NEXT_STEPS["en"])
    if agent_type not in lang_steps:
        agent_type = "custom"
    steps = lang_steps[agent_type]

    add_rule(doc, before_pt=24, after_pt=16, color="FF5119")

    ns_title = doc.add_paragraph()
    ns_title.paragraph_format.space_before = Pt(0)
    ns_title.paragraph_format.space_after  = Pt(10)
    ns_run = ns_title.add_run(L["section_next_steps"].upper())
    set_run_font(ns_run, FONT_TITLE, SIZE_H1, bold=True, color=ORANGE)

    for step in steps:
        sp = doc.add_paragraph()
        sp.paragraph_format.space_before = Pt(4)
        sp.paragraph_format.space_after  = Pt(4)
        sp.paragraph_format.left_indent  = Cm(0.2)
        arrow    = sp.add_run("→  ")
        step_run = sp.add_run(step)
        set_run_font(arrow,    FONT_TITLE, SIZE_BODY, bold=True, color=ORANGE)
        set_run_font(step_run, FONT_BODY,  SIZE_BODY, color=DARK)

    # ══════════════════════════════════════════════════
    # CHECKLIST
    # ══════════════════════════════════════════════════
    add_rule(doc, before_pt=24, after_pt=16, color="FF5119")

    cl_title = doc.add_paragraph()
    cl_title.paragraph_format.space_before = Pt(0)
    cl_title.paragraph_format.space_after  = Pt(10)
    cl_run = cl_title.add_run(L["checklist_title"])
    set_run_font(cl_run, FONT_TITLE, SIZE_H1, bold=True, color=ORANGE)

    for label, detail in L["checklist_items"]:
        cp = doc.add_paragraph()
        cp.paragraph_format.space_before = Pt(4)
        cp.paragraph_format.space_after  = Pt(4)
        cp.paragraph_format.left_indent  = Cm(0.2)

        checkbox   = cp.add_run("☐  ")
        label_run  = cp.add_run(label)
        detail_run = cp.add_run(f"  —  {detail}")
        set_run_font(checkbox,   FONT_TITLE, SIZE_BODY, bold=True, color=ORANGE)
        set_run_font(label_run,  FONT_TITLE, SIZE_BODY, bold=True, color=DARK)
        set_run_font(detail_run, FONT_BODY,  SIZE_BODY, color=GRAY)

    foot = doc.add_paragraph()
    foot.paragraph_format.space_before = Pt(24)
    foot.paragraph_format.space_after  = Pt(0)
    foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    foot_run = foot.add_run(L["footer"].format(date=date_str, version=VERSION))
    set_run_font(foot_run, FONT_TITLE, SIZE_SMALL, color=GRAY)

    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))

    # Export config alongside .docx for reference
    json_out = out.with_suffix(".json")
    with open(json_out, "w", encoding="utf-8") as jf:
        json.dump(config, jf, ensure_ascii=False, indent=2)

    print(L["success"].format(path=str(out)))
    return str(out)


if __name__ == "__main__":
    if len(sys.argv) < 3:
        print("Usage: python generate_docx.py <config.json> <output.docx>")
        sys.exit(1)

    config_path = sys.argv[1]
    output_path = sys.argv[2]

    config_size = Path(config_path).stat().st_size
    if config_size > 1_048_576:  # 1 MB
        print("Error: config file exceeds 1 MB limit", file=sys.stderr)
        sys.exit(1)

    with open(config_path, "r", encoding="utf-8") as f:
        config = json.load(f)

    generate_document(config, output_path)
